from docx import Document
from docx.shared import Pt, RGBColor
from nexus.services.resume_docx_text import _text

NAVY = RGBColor(25, 58, 83)
BLUE = RGBColor(31, 129, 171)
MUTED = RGBColor(91, 110, 124)

def _section(document: Document, title: str):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = BLUE


def _line(document: Document, text: str, *, bold_prefix: str | None = None):
    if not text:
        return
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)


def _bullet(document: Document, text: str):
    if text:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(text)


def _extras(document: Document, item: dict, known: set[str]):
    for key, value in item.items():
        if key not in known and value not in (None, "", [], {}):
            _line(document, f"{key.replace('_', ' ').title()}: {_text(value)}")
