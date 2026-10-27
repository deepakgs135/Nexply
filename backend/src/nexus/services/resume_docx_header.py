from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from nexus.services.resume_docx_components import BLUE, MUTED, NAVY, _text


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.font.color.rgb = MUTED


def render_header(document: Document, data: dict) -> None:
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(2)
    run = heading.add_run(_text(data.get("full_name")) or "Candidate Resume")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = NAVY
    if data.get("professional_title"):
        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(4)
        title_run = title.add_run(_text(data["professional_title"]))
        title_run.font.size = Pt(11)
        title_run.font.color.rgb = BLUE
    contact = [data.get("email"), data.get("phone"), data.get("location"), *data.get("websites", [])]
    contact_text = "  |  ".join(_text(item) for item in contact if _text(item))
    if contact_text:
        line = document.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        line.paragraph_format.space_after = Pt(9)
        line.add_run(contact_text).font.size = Pt(8)
