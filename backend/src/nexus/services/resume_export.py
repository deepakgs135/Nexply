"""Create an in-memory, editable DOCX resume from structured profile data."""

from io import BytesIO
from html.parser import HTMLParser
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from nexus.core.debug import debug_trace
from nexus.core.schemas.resume import ResumeData


NAVY = RGBColor(25, 58, 83)
BLUE = RGBColor(31, 129, 171)
MUTED = RGBColor(91, 110, 124)


class _ResumeTextParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]):
        if tag == "br":
            self.parts.append("\n")
        elif tag == "li":
            self.parts.append("\n• ")
        elif tag == "a":
            href = next((value for name, value in attrs if name == "href"), None)
            if href and href.startswith(("https://", "http://")):
                self.parts.append(f" ({href})")

    def handle_endtag(self, tag: str):
        if tag in {"p", "div", "li"}:
            self.parts.append("\n")

    def handle_data(self, data: str):
        self.parts.append(data)

    def text(self) -> str:
        return re.sub(r"[ \t]*\n[ \t]*", "\n", "".join(self.parts)).strip()


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        items = [_text(item) for item in value]
        return " · ".join(text for text in items if text)
    parser = _ResumeTextParser()
    parser.feed(str(value))
    return parser.text()


def _section(document: Document, title: str):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = BLUE


def _line(document: Document, text: str, *, bold_prefix: str | None = None):
    if not text:
        return
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)


def _bullet(document: Document, text: str):
    if text:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(text)


def _extras(document: Document, item: dict, known: set[str]):
    for key, value in item.items():
        if key not in known and value not in (None, "", [], {}):
            _line(document, f"{key.replace('_', ' ').title()}: {_text(value)}")


@debug_trace
def build_resume_docx(profile: dict) -> bytes:
    """Render validated profile facts into a simple, polished DOCX document."""
    data = ResumeData.model_validate(profile).model_dump(mode="json")
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.font.color.rgb = MUTED

    name = _text(data.get("full_name")) or "Candidate Resume"
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(2)
    run = heading.add_run(name)
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = NAVY
    if data.get("professional_title"):
        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(4)
        title_run = title.add_run(_text(data["professional_title"]))
        title_run.font.size = Pt(11)
        title_run.font.color.rgb = BLUE
    contact = [data.get("email"), data.get("phone"), data.get("location"), *data.get("websites", [])]
    contact_text = "  |  ".join(_text(item) for item in contact if _text(item))
    if contact_text:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(9)
        p.add_run(contact_text).font.size = Pt(8)

    if data.get("professional_summary"):
        _section(document, "Professional Summary")
        _line(document, _text(data["professional_summary"]))

    if data.get("experience"):
        _section(document, "Professional Experience")
        for item in data["experience"]:
            role = " · ".join(value for value in (_text(item.get("job_title")), _text(item.get("company"))) if value)
            _line(document, role, bold_prefix=_text(item.get("job_title")))
            meta = " · ".join(value for value in (_text(item.get("location")), _text(item.get("employment_type")), " – ".join(value for value in (_text(item.get("start_date")), _text(item.get("end_date"))) if value)) if value)
            _line(document, meta)
            _line(document, _text(item.get("description")))
            for field in ("responsibilities", "achievements"):
                for detail in item.get(field, []):
                    _bullet(document, _text(detail))
            if item.get("technologies"):
                _line(document, "Technologies: " + _text(item["technologies"]))
            _extras(document, item, {"company", "job_title", "location", "employment_type", "start_date", "end_date", "description", "responsibilities", "achievements", "technologies"})

    if data.get("skills"):
        _section(document, "Skills")
        _line(document, _text(data["skills"]))

    if data.get("education"):
        _section(document, "Education")
        for item in data["education"]:
            title = " · ".join(value for value in (_text(item.get("degree")), _text(item.get("field_of_study")), _text(item.get("institution"))) if value)
            _line(document, title, bold_prefix=_text(item.get("degree")))
            date_range = " – ".join(value for value in (_text(item.get("start_date")), _text(item.get("end_date"))) if value)
            education_meta = " · ".join(value for value in (_text(item.get("location")), date_range, _text(item.get("cgpa_or_percentage"))) if value)
            _line(document, education_meta)
            for field in ("coursework", "honors"):
                if item.get(field):
                    _line(document, f"{field.replace('_', ' ').title()}: {_text(item[field])}")
            if item.get("graduation_year") is not None:
                _line(document, f"Graduation year: {item['graduation_year']}")
            _extras(document, item, {"degree", "field_of_study", "institution", "location", "start_date", "end_date", "graduation_year", "cgpa_or_percentage", "coursework", "honors"})

    if data.get("projects"):
        _section(document, "Projects")
        for item in data["projects"]:
            title = " · ".join(value for value in (_text(item.get("name")), _text(item.get("role"))) if value)
            _line(document, title, bold_prefix=_text(item.get("name")))
            _line(document, _text(item.get("description")))
            for field in ("responsibilities", "outcomes"):
                for detail in item.get(field, []):
                    _bullet(document, _text(detail))
            if item.get("technologies"):
                _line(document, "Technologies: " + _text(item["technologies"]))
            project_meta = " · ".join(value for value in (_text(item.get("url")), " – ".join(value for value in (_text(item.get("start_date")), _text(item.get("end_date"))) if value)) if value)
            _line(document, project_meta)
            _extras(document, item, {"name", "role", "description", "start_date", "end_date", "url", "technologies", "responsibilities", "outcomes"})

    for section_title, key in (("Certifications", "certifications"), ("Languages", "languages"), ("Awards", "awards"), ("Volunteer Experience", "volunteer_experience"), ("Publications", "publications"), ("Interests", "interests"), ("References", "references")):
        items = data.get(key) or []
        if not items:
            continue
        _section(document, section_title)
        if key in ("volunteer_experience", "publications"):
            for item in items:
                _line(document, " · ".join(value for value in (_text(item.get("job_title") or item.get("title")), _text(item.get("company") or item.get("publisher"))) if value), bold_prefix=_text(item.get("job_title") or item.get("title")))
                _line(document, " · ".join(value for value in (_text(item.get("description")), _text(item.get("publication_date")), _text(item.get("url"))) if value))
                for detail in item.get("responsibilities", []):
                    _bullet(document, _text(detail))
                _extras(document, item, {"company", "job_title", "location", "employment_type", "start_date", "end_date", "description", "responsibilities", "achievements", "technologies", "title", "publisher", "publication_date", "url"})
        else:
            for item in items:
                _bullet(document, _text(item))

    for item in data.get("additional_sections", []):
        _section(document, _text(item.get("title")) or "Additional Information")
        for detail in item.get("details", []):
            _bullet(document, _text(detail))
        _extras(document, item, {"title", "details"})

    output = BytesIO()
    document.save(output)
    return output.getvalue()
