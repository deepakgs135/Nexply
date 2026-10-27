from unittest.mock import Mock
from io import BytesIO

import pytest
from bson import ObjectId
from docx import Document
from fastapi.testclient import TestClient

from nexus.api.main import app
from nexus.api.auth import get_current_user
from nexus.api.routes.resume import extract as extract_route
from nexus.api.routes.resume import history as history_route
from nexus.api.routes.resume import ai as ai_resume_route
from nexus.api.routes.resume import get as get_route
from nexus.api.routes.resume import parse as parse_route
from nexus.api.routes.resume import upload as upload_route
from nexus.api.routes.health import db as health_db

TEST_USER = {
    "id": 7,
    "app_id": "0e0b3552-a55d-49e9-97c2-bcc4a311a4c3",
    "full_name": "Test User",
    "email": "test@example.com",
    "is_active": True,
}


@pytest.fixture
def client():
    app.dependency_overrides[get_current_user] = lambda: TEST_USER
    return TestClient(app)


@pytest.fixture(autouse=True)
def clear_overrides():
    yield
    app.dependency_overrides.clear()


def test_openapi_documents_resume_endpoints(client):
    response = client.get("/openapi.json")

    assert response.status_code == 200
    paths = response.json()["paths"]
    assert "/health" in paths
    assert "/resumes/upload" in paths
    assert "/resumes/{document_id}/parse" in paths
    assert "/resumes/{document_id}/extract" in paths
    assert "/resumes/{document_id}" in paths
    assert "/resumes/latest" in paths
    assert "/resumes/history" not in paths
    assert "/resumes/{document_id}/profile" in paths


def test_root_serves_resume_workspace(client):
    from nexus.api.main import FRONTEND_DIST

    response = client.get("/")

    assert response.status_code == 200
    if (FRONTEND_DIST / "index.html").is_file():
        assert 'id="root"' in response.text
        assert "NEXPLY" in response.text
    else:
        assert response.json()["status"] == "running"


def test_ui_config_exposes_upload_limits(client):
    response = client.get("/api/config")

    assert response.status_code == 200
    assert response.json() == {
        "max_upload_size_mb": 15,
        "supported_formats": [".pdf", ".docx", ".txt"],
    }


def test_ui_build_assets_are_served(client):
    from nexus.api.main import FRONTEND_DIST

    assets = list((FRONTEND_DIST / "assets").glob("*.js"))
    if not assets:
        pytest.skip("React assets have not been built in this checkout.")
    response = client.get(f"/assets/{assets[0].name}")

    assert response.status_code == 200
    assert "Profile" in response.text or "resume" in response.text.lower()


def test_ui_favicon_is_served(client):
    from nexus.api.main import FRONTEND_DIST

    if not (FRONTEND_DIST / "nexply.svg").is_file():
        pytest.skip("React assets have not been built in this checkout.")
    response = client.get("/nexply.svg")

    assert response.status_code == 200
    assert "<svg" in response.text


def test_health_checks_mongodb(client, monkeypatch):
    ping = Mock(return_value={"ok": 1.0})
    monkeypatch.setattr(health_db, "command", ping)

    response = client.get("/health")

    assert response.status_code == 200
    assert response.json() == {"status": "ok", "database": "nexply"}
    ping.assert_called_once_with("ping")


def test_upload_accepts_text_file_and_returns_id(client, monkeypatch):
    save = Mock(return_value="resume-test-id")
    monkeypatch.setattr(upload_route, "save_original_file", save)
    monkeypatch.setattr(upload_route, "retain_latest_resume", Mock(return_value=0))

    response = client.post(
        "/resumes/upload",
        files={"file": ("sample.txt", b"Sample resume text", "text/plain")},
        headers={"Idempotency-Key": "64a195bb-e55c-4b4a-91e0-fcd7623fd010"},
    )

    assert response.status_code == 201
    assert response.json()["document_id"] == "resume-test-id"
    assert response.json()["parsing_status"] == "pending"
    assert save.call_args.kwargs["app_id"] == TEST_USER["app_id"]
    assert save.call_args.kwargs["upload_request_id"] == "64a195bb-e55c-4b4a-91e0-fcd7623fd010"


def test_upload_requires_idempotency_key(client):
    response = client.post(
        "/resumes/upload",
        files={"file": ("sample.txt", b"Sample resume text", "text/plain")},
    )

    assert response.status_code == 422


def test_upload_rejects_unsupported_file(client):
    response = client.post(
        "/resumes/upload",
        files={"file": ("sample.exe", b"not supported", "application/octet-stream")},
        headers={"Idempotency-Key": "de7c49f0-a591-4dcf-8a4d-b4925bdbbf62"},
    )

    assert response.status_code == 415


def test_parse_endpoint_returns_service_result(client, monkeypatch):
    result = {
        "document_id": "resume-test-id",
        "parsing_status": "completed",
        "character_count": 18,
    }
    parse = Mock(return_value=result)
    monkeypatch.setattr(parse_route, "parse_resume", parse)

    response = client.post("/resumes/resume-test-id/parse")

    assert response.status_code == 200
    assert response.json() == result
    parse.assert_called_once_with("resume-test-id", TEST_USER["app_id"])


def test_extract_endpoint_returns_structured_result(client, monkeypatch):
    result = {
        "document_id": "resume-test-id",
        "llm_status": "completed",
        "structured_result": {"full_name": "Sample Person", "skills": []},
    }
    extract = Mock(return_value=result)
    monkeypatch.setattr(extract_route, "extract_resume", extract)

    response = client.post("/resumes/resume-test-id/extract")

    assert response.status_code == 200
    assert response.json() == result
    extract.assert_called_once_with("resume-test-id", TEST_USER["app_id"])


def test_get_endpoint_returns_resume_without_gridfs_id(client, monkeypatch):
    record = {
        "_id": ObjectId("6add939dfffc8e3344105dd8"),
        "original_file_id": ObjectId("6add939dfffc8e3344105dd9"),
        "original_filename": "sample.txt",
        "parsed_text": "Sample resume text",
    }
    get_resume = Mock(return_value=record.copy())
    monkeypatch.setattr(get_route, "get_resume_details", get_resume)

    response = client.get("/resumes/6add939dfffc8e3344105dd8")

    assert response.status_code == 200
    body = response.json()
    assert body["document_id"] == "6add939dfffc8e3344105dd8"
    assert "original_file_id" not in body
    assert body["parsed_text"] == "Sample resume text"
    get_resume.assert_called_once_with("6add939dfffc8e3344105dd8", TEST_USER["app_id"])


def test_latest_resume_returns_only_current_profile(client, monkeypatch):
    record = {
        "_id": ObjectId("6add939dfffc8e3344105dd8"),
        "original_filename": "resume.pdf",
        "original_size_bytes": 12345,
        "created_at": "2026-10-25T12:00:00Z",
        "parsing_status": "completed",
        "llm_status": "completed",
        "structured_result": {"full_name": "Sample Candidate", "professional_title": "Engineer", "experience": [{
            "job_title": "Engineer", "location": None, "location_string": "Pune", "start_date": "01/2024", "end_date": "Present",
            "responsibilities": [], "responsibilities_and_achievements": ["Built complete feature"],
        }]},
        "parsed_text": "Engineer\n01/2024 – Present\nPune",
    }
    record["app_id"] = TEST_USER["app_id"]
    record["parsed_text"] = "must not be returned"
    get_latest = Mock(return_value=record.copy())
    monkeypatch.setattr(history_route, "get_latest_resume_for_owner", get_latest)
    monkeypatch.setattr(history_route.postgres, "get_candidate_profile_source_document_id", Mock(return_value=None))
    sync_profile = Mock()
    monkeypatch.setattr(history_route, "sync_candidate_profile_to_sql", sync_profile)

    response = client.get("/resumes/latest")

    assert response.status_code == 200
    body = response.json()["resume"]
    assert body["document_id"] == "6add939dfffc8e3344105dd8"
    assert body["structured_result"]["full_name"] == "Sample Candidate"
    assert body["structured_result"]["experience"][0]["location"] == "Pune"
    assert body["structured_result"]["experience"][0]["description"] == "<ul><li>Built complete feature</li></ul>"
    assert body["structured_result"]["experience"][0]["responsibilities"] == []
    assert "app_id" not in body
    assert "parsed_text" not in body
    get_latest.assert_called_once_with(TEST_USER["app_id"])
    sync_profile.assert_called_once_with(
        "6add939dfffc8e3344105dd8", TEST_USER["app_id"], "resume.pdf", body["structured_result"],
    )


def test_resume_history_profile_is_scoped_to_authenticated_owner(client, monkeypatch):
    record = {
        "_id": ObjectId("6add939dfffc8e3344105dd8"),
        "app_id": TEST_USER["app_id"],
        "structured_result": {"full_name": "Sample Candidate", "skills": ["Python"]},
        "parsed_text": "private extracted source text",
    }
    get_record = Mock(return_value=record.copy())
    monkeypatch.setattr(history_route, "get_resume_details", get_record)

    response = client.get("/resumes/6add939dfffc8e3344105dd8/profile")

    assert response.status_code == 200
    assert response.json()["structured_result"]["skills"] == ["Python"]
    assert "parsed_text" not in response.json()
    get_record.assert_called_once_with("6add939dfffc8e3344105dd8", TEST_USER["app_id"])


def test_structured_profile_edit_updates_existing_owner_record(client, monkeypatch):
    update = Mock(return_value=True)
    monkeypatch.setattr(history_route, "update_resume", update)
    monkeypatch.setattr(history_route, "get_resume_details", Mock(return_value={"parsed_text": ""}))
    sync_profile = Mock()
    monkeypatch.setattr(history_route, "sync_candidate_profile_to_sql", sync_profile)
    profile = {"full_name": "Corrected Name", "websites": ["https://example.com"], "skills": ["Python"], "experience": [{
        "job_title": "Engineer", "location": None, "location_string": "Chennai",
        "responsibilities": [], "responsibilities_and_achievements": ["Built a system", "Improved accuracy"],
    }]}

    response = client.put("/resumes/6add939dfffc8e3344105dd8/structured-result", json={"structured_result": profile})

    assert response.status_code == 200
    assert response.json()["structured_result"]["full_name"] == "Corrected Name"
    experience = response.json()["structured_result"]["experience"][0]
    assert experience["location"] == "Chennai"
    assert experience["description"] == "<ul><li>Built a system</li><li>Improved accuracy</li></ul>"
    assert experience["responsibilities"] == []
    assert "location_string" not in experience
    update.assert_called_once()
    args = update.call_args.args
    assert args[0] == "6add939dfffc8e3344105dd8"
    assert args[1] == TEST_USER["app_id"]
    assert args[2]["structured_result"]["websites"] == ["https://example.com"]
    assert args[2]["structured_result"]["experience"][0]["description"] == experience["description"]
    sync_profile.assert_called_once()
    assert sync_profile.call_args.args[0:3] == (
        "6add939dfffc8e3344105dd8", TEST_USER["app_id"], None,
    )


def test_structured_profile_edit_returns_not_found_for_non_owner(client, monkeypatch):
    update = Mock(return_value=False)
    monkeypatch.setattr(history_route, "update_resume", update)
    monkeypatch.setattr(history_route, "get_resume_details", Mock(return_value={"parsed_text": ""}))

    response = client.put("/resumes/6add939dfffc8e3344105dd8/structured-result", json={"structured_result": {"full_name": "Other"}})

    assert response.status_code == 404
    update.assert_called_once()
    assert update.call_args.args[1] == TEST_USER["app_id"]


def test_job_preferences_update_is_scoped_to_authenticated_app_id(client, monkeypatch):
    upsert = Mock(return_value={"app_id": TEST_USER["app_id"]})
    monkeypatch.setattr(history_route.postgres, "upsert_job_preferences", upsert)
    preferences = {
        "current_ctc": "8 LPA", "expected_ctc": "12 LPA", "expected_ctc_range": "10-14 LPA",
        "minimum_expected_ctc": "10 LPA", "currency": "INR", "preferred_locations": ["Pune", "Remote"],
    }

    response = client.put("/resumes/preferences", json={"job_preferences": preferences})

    assert response.status_code == 200
    assert response.json()["job_preferences"]["preferred_locations"] == ["Pune", "Remote"]
    upsert.assert_called_once()
    assert upsert.call_args.args[0] == TEST_USER["app_id"]
    assert upsert.call_args.args[1]["minimum_expected_ctc"] == "10 LPA"


def test_job_preferences_load_uses_authenticated_app_id(client, monkeypatch):
    load = Mock(return_value={"current_ctc": "8 LPA", "preferred_locations": ["Pune"]})
    monkeypatch.setattr(history_route.postgres, "get_job_preferences", load)

    response = client.get("/resumes/preferences")

    assert response.status_code == 200
    assert response.json()["configured"] is True
    assert response.json()["job_preferences"]["preferred_locations"] == ["Pune"]
    load.assert_called_once_with(TEST_USER["app_id"])


def test_job_preferences_update_returns_setup_instructions_when_table_missing(client, monkeypatch):
    from psycopg.errors import UndefinedTable

    upsert = Mock(side_effect=UndefinedTable("missing preferences table"))
    monkeypatch.setattr(history_route.postgres, "upsert_job_preferences", upsert)

    response = client.put("/resumes/preferences", json={"job_preferences": {}})

    assert response.status_code == 503
    assert "docs/sql/user_job_preferences.sql" in response.json()["detail"]
    assert upsert.call_args.args[0] == TEST_USER["app_id"]


def test_structured_profile_edit_rejects_invalid_data(client, monkeypatch):
    update = Mock(return_value=True)
    monkeypatch.setattr(history_route, "update_resume", update)
    monkeypatch.setattr(history_route, "get_resume_details", Mock(return_value={"parsed_text": ""}))

    response = client.put("/resumes/6add939dfffc8e3344105dd8/structured-result", json={"structured_result": {"education": [{"graduation_year": "not a year"}]}})

    assert response.status_code == 422
    update.assert_not_called()


def test_ai_smart_uses_owner_profile_and_does_not_persist_draft(client, monkeypatch):
    source = {"structured_result": {"full_name": "Sample Candidate", "skills": ["Python"]}}
    generated = {"full_name": "Sample Candidate", "professional_summary": "A clearer summary.", "skills": ["Python"]}
    get_record = Mock(return_value=source)
    models = Mock(return_value=["llama3.1:8b", "qwen3.5:9b"])
    writer = Mock(return_value=generated)
    monkeypatch.setattr(ai_resume_route, "get_resume_details", get_record)
    monkeypatch.setattr(ai_resume_route, "available_resume_models", models)
    monkeypatch.setattr(ai_resume_route, "write_resume", writer)

    response = client.post("/resumes/6add939dfffc8e3344105dd8/ai-resume", json={"model": "llama3.1:8b", "target_job_description": "Data Engineer"})

    assert response.status_code == 200
    assert response.json()["structured_result"]["professional_summary"] == "A clearer summary."
    get_record.assert_called_once_with("6add939dfffc8e3344105dd8", TEST_USER["app_id"])
    writer.assert_called_once_with(source["structured_result"], "llama3.1:8b", "Data Engineer")


def test_ai_smart_cannot_access_another_users_resume(client, monkeypatch):
    get_record = Mock(return_value=None)
    writer = Mock()
    monkeypatch.setattr(ai_resume_route, "get_resume_details", get_record)
    monkeypatch.setattr(ai_resume_route, "write_resume", writer)

    response = client.post("/resumes/6add939dfffc8e3344105dd8/ai-resume", json={})

    assert response.status_code == 404
    get_record.assert_called_once_with("6add939dfffc8e3344105dd8", TEST_USER["app_id"])
    writer.assert_not_called()


def test_rex_rewrites_only_requested_field_from_owner_profile(client, monkeypatch):
    source = {"structured_result": {"full_name": "Sample Candidate", "professional_summary": "Old summary."}}
    get_record = Mock(return_value=source)
    rewrite = Mock(return_value={"model": "llama3.1:8b", "rewritten_text": "Improved factual summary."})
    monkeypatch.setattr(ai_resume_route, "get_resume_details", get_record)
    monkeypatch.setattr(ai_resume_route, "rewrite_resume_text", rewrite)

    response = client.post("/resumes/6add939dfffc8e3344105dd8/rex/rewrite", json={"field_path": ["professional_summary"]})

    assert response.status_code == 200
    assert response.json()["rewritten_text"] == "Improved factual summary."
    get_record.assert_called_once_with("6add939dfffc8e3344105dd8", TEST_USER["app_id"])
    rewrite.assert_called_once_with(source["structured_result"], ["professional_summary"], None, "brief")


def test_rex_manual_profile_interview_returns_dynamic_profile(client, monkeypatch):
    profile = {"full_name": "Deepak Kumar", "experience": [], "education": [], "skills": []}
    interview = Mock(return_value={
        "model": "llama3.1:8b", "assistant_message": "What is your current designation at Tata Technologies?",
        "ready_to_save": True, "structured_result": profile,
    })
    monkeypatch.setattr(ai_resume_route, "interview_resume_profile", interview)

    response = client.post("/resumes/manual-profile/interview", json={
        "messages": [{"role": "user", "content": "I am Deepak Kumar and work at Tata Technologies."}],
    })

    assert response.status_code == 200
    assert "current designation" in response.json()["assistant_message"]
    interview.assert_called_once()


def test_manual_profile_saves_owner_scoped_latest_record(client, monkeypatch):
    profile = {"full_name": "Deepak Kumar", "professional_title": "Senior Engineer"}
    create = Mock()
    upload = Mock(return_value=ObjectId("6add939dfffc8e3344105dd9"))
    monkeypatch.setattr(ai_resume_route, "create_resume", create)
    monkeypatch.setattr(ai_resume_route.file_bucket, "upload_from_stream", upload)
    monkeypatch.setattr(ai_resume_route, "retain_latest_resume", Mock(return_value=0))
    sync_profile = Mock()
    monkeypatch.setattr(ai_resume_route, "sync_candidate_profile_to_sql", sync_profile)

    response = client.post("/resumes/manual-profile", json={"structured_result": profile})

    assert response.status_code == 200
    assert response.json()["structured_result"]["full_name"] == "Deepak Kumar"
    saved_record, owner_id = create.call_args.args
    assert saved_record["source"] == "rex_manual_profile"
    assert saved_record["original_filename"] == "Deepak-Kumar-resume.docx"
    assert saved_record["original_content_type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert saved_record["original_size_bytes"] > 0
    assert saved_record["original_file_id"] == upload.return_value
    assert saved_record["_id"] == ObjectId(response.json()["document_id"])
    assert saved_record["structured_result"]["full_name"] == "Deepak Kumar"
    assert owner_id == TEST_USER["app_id"]
    filename, content, metadata = upload.call_args.args[0], upload.call_args.args[1], upload.call_args.kwargs["metadata"]
    assert filename == "Deepak-Kumar-resume.docx"
    assert content.startswith(b"PK")
    generated_doc = Document(BytesIO(content))
    assert any("Senior Engineer" in paragraph.text for paragraph in generated_doc.paragraphs)
    assert metadata["stage"] == "manual_resume"
    assert metadata["document_id"] == response.json()["document_id"]
    assert metadata["app_id"] == TEST_USER["app_id"]
    sync_profile.assert_called_once_with(
        response.json()["document_id"], TEST_USER["app_id"], filename, response.json()["structured_result"],
    )


def test_rex_manual_profile_cleans_gridfs_file_when_record_save_fails(client, monkeypatch):
    monkeypatch.setattr(ai_resume_route, "create_resume", Mock(side_effect=RuntimeError("database unavailable")))
    upload = Mock(return_value=ObjectId("6add939dfffc8e3344105dd9"))
    delete = Mock()
    monkeypatch.setattr(ai_resume_route.file_bucket, "upload_from_stream", upload)
    monkeypatch.setattr(ai_resume_route.file_bucket, "delete", delete)

    response = client.post("/resumes/manual-profile", json={"structured_result": {"full_name": "Deepak Kumar"}})

    assert response.status_code == 500
    assert response.json()["detail"] == "Unable to save the manually created resume."
    delete.assert_called_once_with(upload.return_value)


def test_ai_resume_export_returns_docx_for_owned_resume(client, monkeypatch):
    monkeypatch.setattr(ai_resume_route, "get_resume_details", Mock(return_value={"structured_result": {"full_name": "Sample Candidate"}}))
    response = client.post("/resumes/6add939dfffc8e3344105dd8/ai-resume/export", json={"structured_result": {"full_name": "Sample Candidate", "skills": ["Python"]}})

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert response.content.startswith(b"PK")


def test_resume_export_returns_docx_without_existing_resume(client):
    response = client.post("/resumes/export", json={"structured_result": {
        "full_name": "Sample Candidate",
        "professional_summary": "<p>Led <strong>platform delivery</strong>.</p>",
        "skills": ["Python"],
    }})

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert response.headers["content-disposition"].endswith('filename="Sample-Candidate-resume.docx"')
    assert response.content.startswith(b"PK")
    generated_doc = Document(BytesIO(response.content))
    paragraphs = "\n".join(paragraph.text for paragraph in generated_doc.paragraphs)
    assert "Led platform delivery." in paragraphs
    assert "<strong>" not in paragraphs


def test_resume_export_requires_authentication():
    app.dependency_overrides.clear()
    unauthenticated_client = TestClient(app)

    response = unauthenticated_client.post("/resumes/export", json={"structured_result": {"full_name": "Sample Candidate"}})

    assert response.status_code == 401


def test_resume_delete_removes_authenticated_users_item(client, monkeypatch):
    delete = Mock(return_value=True)
    monkeypatch.setattr(history_route, "delete_resume", delete)

    response = client.delete("/resumes/6add939dfffc8e3344105dd8")

    assert response.status_code == 204
    delete.assert_called_once_with("6add939dfffc8e3344105dd8", TEST_USER["app_id"])


def test_resume_delete_returns_not_found_when_record_is_not_owned(client, monkeypatch):
    delete = Mock(return_value=False)
    monkeypatch.setattr(history_route, "delete_resume", delete)

    response = client.delete("/resumes/6add939dfffc8e3344105dd8")

    assert response.status_code == 404


def test_resume_delete_endpoint_requires_authentication(monkeypatch):
    app.dependency_overrides.clear()
    unauthenticated_client = TestClient(app)
    delete = Mock(return_value=True)
    monkeypatch.setattr(history_route, "delete_resume", delete)

    response = unauthenticated_client.delete("/resumes/6add939dfffc8e3344105dd8")

    assert response.status_code == 401
    delete.assert_not_called()


def test_resume_routes_require_authentication():
    app.dependency_overrides.clear()
    unauthenticated_client = TestClient(app)

    response = unauthenticated_client.get("/resumes/6add939dfffc8e3344105dd8")

    assert response.status_code == 401


@pytest.mark.parametrize("path,method", [
    ("/resumes/upload", "post"),
    ("/resumes/6add939dfffc8e3344105dd8/parse", "post"),
    ("/resumes/6add939dfffc8e3344105dd8/extract", "post"),
])
def test_resume_write_and_processing_routes_require_authentication(path, method):
    app.dependency_overrides.clear()
    unauthenticated_client = TestClient(app)
    kwargs = {"files": {"file": ("resume.txt", b"candidate", "text/plain")}} if path.endswith("upload") else {}

    response = getattr(unauthenticated_client, method)(path, **kwargs)

    assert response.status_code == 401


def test_openapi_exposes_auth_routes(client):
    response = client.get("/openapi.json")

    assert response.status_code == 200
    paths = response.json()["paths"]
    assert "/auth/register" in paths
    assert "/auth/login" in paths
    assert "/auth/me" in paths
